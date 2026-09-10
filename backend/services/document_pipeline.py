from __future__ import annotations

import io
import os
import uuid
from pathlib import Path

from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from models import LibraryDocument, LibraryDocumentChunk, ProcessingStatus
from services.accession import apply_processing_outcome, mark_processing_fields
from services.chunking import ParsedChunk, split_with_overlap
from services.embeddings import EmbeddingService

WORKSPACE_ROOT = Path(__file__).resolve().parent.parent.parent
MAX_EXTRACT_CHARS = 400_000


def documents_dir() -> Path:
    configured = os.getenv("DOCUMENTS_DIR", "data/documents")
    path = Path(configured)
    if not path.is_absolute():
        path = WORKSPACE_ROOT / path
    path.mkdir(parents=True, exist_ok=True)
    return path


def document_file_path(document_id: uuid.UUID, filename: str) -> Path:
    suffix = Path(filename).suffix.lower() or ".bin"
    return documents_dir() / f"{document_id}{suffix}"


class ParsedDocument:
    def __init__(
        self,
        title: str,
        extracted_text: str,
        chunks: list[ParsedChunk],
    ) -> None:
        self.title = title
        self.extracted_text = extracted_text
        self.chunks = chunks


class DocumentPipeline:
    def __init__(self, embeddings: EmbeddingService):
        self.embeddings = embeddings

    def parse(self, file_path: str, fallback_title: str) -> ParsedDocument:
        path = Path(file_path)
        content = path.read_bytes()
        text = extract_document_text(path.name, content)
        title = fallback_title or path.stem or "Untitled document"
        first_line = next(
            (line.strip() for line in text.splitlines() if line.strip()),
            "",
        )
        if first_line and len(first_line) <= 120 and first_line != title:
            title = first_line.lstrip("# ").strip() or title
        chunks = [
            ParsedChunk(text=part, page=None, section=None)
            for part in split_with_overlap(text)
        ]
        if not chunks and text.strip():
            chunks = [ParsedChunk(text=text.strip()[:1800], page=None, section=None)]
        return ParsedDocument(
            title=title, extracted_text=text[:MAX_EXTRACT_CHARS], chunks=chunks
        )

    def embed_chunks(self, parsed: ParsedDocument) -> list[list[float]]:
        return [self.embeddings.embed_document(chunk.text) for chunk in parsed.chunks]


def extract_document_text(filename: str, content: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        from services.document_classifier import preview_text

        try:
            reader = PdfReader(io.BytesIO(content))
        except Exception:
            return preview_text(content)
        parts: list[str] = []
        for page in reader.pages:
            raw = page.extract_text() or ""
            cleaned = "\n".join(line.strip() for line in raw.splitlines())
            if cleaned.strip():
                parts.append(cleaned)
        return "\n\n".join(parts).strip()[:MAX_EXTRACT_CHARS]
    if suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise RuntimeError("python-docx is required to read .docx files") from exc
        document = Document(io.BytesIO(content))
        parts = [para.text.strip() for para in document.paragraphs if para.text.strip()]
        return "\n".join(parts)[:MAX_EXTRACT_CHARS]
    return content.decode("utf-8", errors="replace").strip()[:MAX_EXTRACT_CHARS]


async def apply_parse_result(
    session: AsyncSession,
    document: LibraryDocument,
    parsed: ParsedDocument,
    embeddings: list[list[float]],
) -> LibraryDocument:
    await session.execute(
        delete(LibraryDocumentChunk).where(
            LibraryDocumentChunk.document_id == document.id
        )
    )
    document.title = parsed.title or document.title
    document.extracted_text = parsed.extracted_text
    apply_processing_outcome(document, chunk_count=len(parsed.chunks))

    for index, (chunk, embedding) in enumerate(
        zip(parsed.chunks, embeddings, strict=True)
    ):
        session.add(
            LibraryDocumentChunk(
                document_id=document.id,
                chunk_index=index,
                text=chunk.text,
                page=chunk.page,
                section=chunk.section,
                embedding=embedding,
                extra={},
            )
        )
    await session.commit()
    await session.refresh(document)
    return document


async def mark_document_status(
    session: AsyncSession,
    document_id: uuid.UUID,
    status: ProcessingStatus,
    error: str | None = None,
) -> None:
    document = await session.get(LibraryDocument, document_id)
    if document is None:
        return
    mark_processing_fields(document, status.value, error)
    await session.commit()


async def chunk_count_for(session: AsyncSession, document_id: uuid.UUID) -> int:
    result = await session.execute(
        select(func.count(LibraryDocumentChunk.id)).where(
            LibraryDocumentChunk.document_id == document_id
        )
    )
    return int(result.scalar_one())
